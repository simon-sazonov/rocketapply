"""The cover letter as a PDF and a DOCX, on the CV's letterhead.

Until 29 Sep 2026 the letter existed only as `.txt`. That is right for a paste
box and wrong for the upload field most applications have, where a text file
looks unfinished next to a typeset CV. So the letter is set in the CV's own
fonts, under the CV's own header — drawn by the same `draw_header`, so the two
documents a recruiter opens side by side cannot drift apart.

Nothing here changes a word. The text is the gated letter exactly as written
to the `.txt`; this module only lays it out. A letter the gates marked
DO_NOT_SEND is never rendered (the caller checks), because a typeset PDF of an
unsupported claim is precisely the file that gets sent by mistake.
"""

from __future__ import annotations

import datetime as dt
import pathlib
from dataclasses import dataclass, field

from .layout import Frame, escape_markup, paragraph_style, register_fonts
from .models import Profile
from .render_pdf import _Canvas, draw_header

__all__ = ["LetterRender", "letter_blocks", "render_letter_pdf", "render_letter_docx",
           "render_letter_files", "LETTER_PDF_SUFFIX", "LETTER_DOCX_SUFFIX"]

LETTER_PDF_SUFFIX = "_Cover_Letter.pdf"
LETTER_DOCX_SUFFIX = "_Cover_Letter.docx"


@dataclass
class LetterRender:
    pdf: pathlib.Path | None = None
    docx: pathlib.Path | None = None
    page_count: int = 0
    #: The header and contact strings, for the test that holds them equal to the CV's.
    sections: dict[str, list[str]] = field(default_factory=dict)
    warnings: list[str] = field(default_factory=list)


def letter_blocks(text: str) -> list[list[str]]:
    """Paragraphs, each a list of its lines. A blank line separates paragraphs;
    a single newline inside one (the sign-off: "Kind regards,\\nName") is kept."""
    blocks: list[list[str]] = []
    for chunk in (text or "").replace("\r\n", "\n").split("\n\n"):
        lines = [line.rstrip() for line in chunk.strip("\n").split("\n")]
        lines = [line for line in lines if line.strip()]
        if lines:
            blocks.append(lines)
    return blocks


def _today() -> str:
    return dt.date.today().strftime("%-d %B %Y")


def render_letter_pdf(profile: Profile, frame: Frame, text: str, path: pathlib.Path | str,
                      *, company: str = "", title: str = "", date: str | None = None
                      ) -> LetterRender:
    from reportlab.platypus import Paragraph

    register_fonts()
    path = pathlib.Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    canvas = _Canvas(path, frame)
    canvas.c.setTitle(f"{profile.personal.name} — Cover letter")
    out = LetterRender(pdf=path)
    f = frame
    geom = f.raw.get("geometry") or {}
    section_gap = float(geom.get("section_gap", 3.5))
    gap = float(geom.get("rule_gap", 5.3))
    width = f.content_width
    bottom = f.page_height - f.margin_bottom
    pages = 1

    y = draw_header(canvas, profile, f)
    # The date on the Email line, flush right — the right of the name line is
    # the title, which the letterhead shares with the CV. It used to sit on the
    # company line, and a long role title ran straight into it.
    canvas.section = "date"
    canvas.right_text(date or _today(), f.body, f.page_width - f.margin_right,
                      canvas.contact_top)
    y += section_gap
    canvas.rule(y)
    y += gap + f.body.line_height()

    style = paragraph_style(f.body.resolved(), f.body.size, f.body.line_height())
    # Company, then role, a line each: no separator to misread, and a role of
    # any length wraps inside the column instead of overprinting.
    canvas.section = "recipient"
    if company:
        canvas.text(company, f.body_bold, f.content_x, y)
        y += f.body.line_height()
    if title:
        y += canvas.para(title, f.body, f.content_x, y, width)
    if company or title:
        y += f.body.line_height()

    canvas.section = "letter"
    spacing = f.body.line_height() * 0.8
    for block in letter_blocks(text):
        # One Paragraph per block; inner newlines become line breaks.
        markup = "<br/>".join(escape_markup(line) for line in block)
        para = Paragraph(markup, style)
        _, height = para.wrapOn(canvas.c, width, 10_000)
        if y + height > bottom:
            canvas.c.showPage()
            pages += 1
            y = f.margin_top
        para.drawOn(canvas.c, f.content_x, canvas._flip(y, height))
        canvas._record("\n".join(block))
        y += height + spacing

    canvas.save()
    out.page_count = pages
    out.sections = {k: list(v) for k, v in canvas.sections.items()}
    if pages > 1:
        out.warnings.append(f"the letter runs to {pages} pages")
    return out


def render_letter_docx(profile: Profile, text: str, path: pathlib.Path | str, *,
                       company: str = "", title: str = "", date: str | None = None
                       ) -> pathlib.Path:
    from docx import Document
    from docx.shared import Pt

    from .render_docx import _style

    path = pathlib.Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    _style(document)
    document.add_heading(f"{profile.personal.name} — {profile.personal.title}", level=0)
    contact = " | ".join(
        v for v in (profile.personal.email, profile.personal.github, profile.personal.linkedin)
        if v
    )
    document.add_paragraph(contact)
    document.add_paragraph(date or _today())
    if company or title:
        header = document.add_paragraph()
        if company:
            header.add_run(company).bold = True
        if company and title:
            header.add_run().add_break()
        if title:
            header.add_run(title)
    for block in letter_blocks(text):
        para = document.add_paragraph()
        for i, line in enumerate(block):
            if i:
                para.add_run().add_break()
            para.add_run(line)
        para.paragraph_format.space_after = Pt(8)
    document.save(str(path))
    return path


def render_letter_files(directory: pathlib.Path | str, stem: str, text: str, *,
                        profile: Profile, frame: Frame, company: str = "", title: str = "",
                        date: str | None = None) -> LetterRender:
    """Both files beside the CV: `<stem>_Cover_Letter.pdf` and `.docx`."""
    directory = pathlib.Path(directory)
    out = render_letter_pdf(profile, frame, text, directory / f"{stem}{LETTER_PDF_SUFFIX}",
                            company=company, title=title, date=date)
    out.docx = render_letter_docx(profile, text, directory / f"{stem}{LETTER_DOCX_SUFFIX}",
                                  company=company, title=title, date=date)
    return out
